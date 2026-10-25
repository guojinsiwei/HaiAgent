from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '长程Agent基准筛选简表.docx'

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(4)
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = Pt(16)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Title'].paragraph_format.space_after = Pt(9)
doc.styles['Subtitle'].font.size = Pt(10)
doc.styles['Subtitle'].font.italic = False
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)

def text(label, content):
    p = doc.add_paragraph()
    p.add_run(label + '：').bold = True
    p.add_run(content)
    return p

def link(label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run(label + '：').font.size = Pt(10)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79'); props.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '19'); props.append(size)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    r.append(props)
    t = OxmlElement('w:t'); t.text = url; r.append(t)
    h.append(r); p._p.append(h)

def entry(title, task, fit, note, links):
    doc.add_heading(title, level=2)
    text('任务内容', task)
    text('适合检验', fit)
    text('选择依据与注意事项', note)
    for label, url in links:
        link(label, url)

doc.add_paragraph('长程 Agent 基准筛选', 'Title')
doc.add_paragraph('供选题与实验讨论使用    核查日期 2026年10月10日', 'Subtitle')
doc.add_paragraph('建议以 NL2Repo-Bench 和 Terminal-Bench 作为主评测，用 UltraHorizon 补充验证记忆与探索机制。以下列出六个候选的任务内容和官方入口，推荐顺序综合考虑任务契合度、认可度与接入成本。')
doc.add_heading('优先推荐', level=1)

entry('NL2Repo Bench',
      '根据自然语言需求，从空目录构建完整、可安装运行的 Python 软件库，包含架构设计、依赖管理、多文件实现和功能交付，共104项任务。',
      '子目标管理、长规格记忆、跨文件一致性，以及避免遗漏功能和提前宣称完成。',
      '已发表于 ICML 2026；任务规格平均约18.8K tokens。原论文中 DeepSeek-V3.2 平均254.3个交互轮次，不等同于API请求数。须防止直接获取目标项目原始源码。',
      [('官方代码', 'https://github.com/multimodal-art-projection/NL2RepoBench'),
       ('正式论文', 'https://proceedings.mlr.press/v306/ding26j.html')])

entry('Terminal Bench',
      '在隔离的终端环境完成真实工作任务，覆盖编程、系统操作、数据处理等；使用任务对应的验证器评测交付结果。',
      '现有 CLI Agent 的实际解题能力、失败恢复、重复调用控制和工具使用效率。',
      '已被主流模型厂商采用。官网当前为4.0，正式评测需锁定版本；并非每题都很长。4.0的8小时是任务超时上限，不是必须跑满的时长。',
      [('官方主页', 'https://www.tbench.ai/'),
       ('版本与任务说明', 'https://www.tbench.ai/news/terminal-bench-4-0')])

entry('UltraHorizon',
      '在 Mystery Grid、Sequence Exploration 和 Alien Genetics Laboratory 三种部分可观测环境中持续试验，归纳隐藏规则。',
      '失败记忆是否被利用、能否跳出固定假设，以及探索是否带来新证据。',
      '2025年发布，专门面向长程探索，适合做机制实验；社区传播范围有限，不建议独自承担主评测。使用文本工具接口，需固定环境种子与评分协议。',
      [('官方代码', 'https://github.com/StarDewXXX/UltraHorizon'),
       ('论文', 'https://arxiv.org/abs/2509.21766')])

doc.add_page_break()
doc.add_heading('备选基准', level=1)

entry('AgencyBench v2',
      '包含32个场景、138个子任务，覆盖游戏、前端、后端、代码、研究和MCP工作流；部分场景需要连续完成多个相互关联的子任务。',
      '子目标衔接、需求与反馈记忆，以及跨阶段修改和成果核验。',
      '官方标注 ACL 2026 Main。建议先考虑后端与代码场景；部分任务需要视觉评委或外部服务。公开的约1M tokens是场景累计消耗，不能当作单次输入长度。',
      [('官方代码', 'https://github.com/GAIR-NLP/AgencyBench'),
       ('论文', 'https://arxiv.org/abs/2601.11044')])

entry('OSWorld 2.0',
      '108个真实电脑操作长工作流，涉及桌面软件、网页和文件等，需要在持续变化的环境中完成完整任务。',
      '长程约束保持、跨来源记忆、隐藏状态推断和最终结果验证。',
      '2026年发布，承接 OSWorld 系列；任务长度突出，但需接入GUI操作与视觉感知。官方当前推荐运行版本为osworld-v2.1，代码、任务、资产和网站版本必须配套。',
      [('官方代码', 'https://github.com/xlang-ai/OSWorld-V2'),
       ('论文', 'https://arxiv.org/abs/2606.29537')])

entry('PaperBench',
      '从零复现20篇 ICML 2024 Spotlight 或 Oral 论文，包括理解方法、编写代码、运行实验和复现实验结果。',
      '科研子目标规划、实验记忆、长流程执行和成果真实性核验。',
      'OpenAI于2025年发布；按分解评分标准评价复现程度。任务计算资源与模型评委成本较高，当前宜作为后续扩展，不优先用于低成本对照。',
      [('官方介绍', 'https://openai.com/index/paperbench/'),
       ('官方代码', 'https://github.com/openai/frontier-evals/tree/main/project/paperbench')])

for paragraph in doc.paragraphs:
    for border in list(paragraph._p.xpath('.//w:pBdr')):
        border.getparent().remove(border)

doc.core_properties.title = '长程 Agent 基准筛选'
doc.core_properties.subject = 'Benchmark任务内容与官方链接'
doc.core_properties.author = ''
doc.core_properties.keywords = 'Agent, Benchmark, 长程任务, 记忆, 探索'
doc.save(OUT)
print(OUT)
print('Hyperlinks:', len([r for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK]))
